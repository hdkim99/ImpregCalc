"""Human-readable and machine-readable recipe exports."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from impregcalc.core.models import RecipeResult


def format_markdown(result: RecipeResult) -> str:
    """Format a complete laboratory-facing recipe with conventions and caveats."""

    request = result.request
    lines = [
        "# ImpregCalc recipe",
        "",
        "## Support and liquid plan",
        "",
        f"- Support: {request.support_name}",
        f"- Dry support mass: {request.support_mass_g:.6g} g",
        f"- Method: {request.method.value}",
        f"- Pore volume: {request.pore_volume_ml_g:.6g} mL g⁻¹"
        if request.pore_volume_ml_g is not None
        else "- Pore volume: not supplied",
        f"- Nominal pore-volume product: {result.nominal_pore_volume_ml:.6g} mL"
        if result.nominal_pore_volume_ml is not None
        else "- Nominal pore-volume product: not calculated",
        f"- Selected total solution volume: {result.selected_solution_volume_ml:.6g} mL",
        f"- Arithmetic impregnation cycles: {request.cycles}",
        f"- Final formulated mass (support + named retained components): "
        f"{result.final_formulated_mass_g:.6g} g",
        "",
        "## Targets and precursor quantities",
        "",
    ]
    for row in result.requirements:
        lines.extend(
            [
                f"### {row.name}",
                "",
                f"- Target: {row.basis_description}",
                f"- Retained component: {row.active_formula}; {row.active_component_mass_g:.6g} g",
                f"- Target-element mass ({row.target_element}): {row.target_element_mass_g:.6g} g",
                f"- Precursor: {row.precursor_formula}",
                f"- Precursor molar mass: {row.precursor_molar_mass_g_mol:.6g} g mol⁻¹",
                f"- Assay/purity: {100 * row.precursor_purity_fraction:.6g}%",
                f"- Required pure precursor equivalent: {row.pure_precursor_mass_g:.6g} g",
                f"- Required commercial precursor to weigh: "
                f"{row.commercial_precursor_mass_g:.6g} g",
                f"- Pure precursor amount: {row.precursor_amount_mol:.6g} mol",
                f"- Required commercial mass concentration: "
                f"{row.concentration_commercial_g_ml:.6g} g mL⁻¹",
                f"- Pure precursor molarity: {row.concentration_pure_mol_l:.6g} mol L⁻¹",
            ]
        )
        if row.atomic_ratio_note:
            lines.append(f"- Ratio note: {row.atomic_ratio_note}")
        if len(row.cycle_portions) > 1:
            lines.extend(
                [
                    "",
                    "| Cycle | Fraction | Precursor / g | Solution / mL |",
                    "|---:|---:|---:|---:|",
                ]
            )
            lines.extend(
                f"| {cycle.cycle} | {cycle.fraction:.6g} | "
                f"{cycle.precursor_mass_g:.6g} | {cycle.solution_volume_ml:.6g} |"
                for cycle in row.cycle_portions
            )
        lines.append("")
    lines.extend(["## Assumptions", ""])
    lines.extend(f"- {item}" for item in result.assumptions)
    lines.extend(["", "## Warnings and scope limits", ""])
    lines.extend(f"- {item}" for item in result.warnings)
    lines.extend(
        [
            "",
            f"Mass-closure residual: {result.mass_closure_residual_g:.3e} g.",
            "",
        ]
    )
    return "\n".join(lines)


def _csv_rows(result: RecipeResult) -> list[dict[str, Any]]:
    return [
        {
            "name": row.name,
            "kind": row.kind.value,
            "basis": row.basis_description,
            "active_formula": row.active_formula,
            "active_component_mass_g": row.active_component_mass_g,
            "target_element": row.target_element,
            "target_element_mass_g": row.target_element_mass_g,
            "precursor_formula": row.precursor_formula,
            "precursor_molar_mass_g_mol": row.precursor_molar_mass_g_mol,
            "precursor_purity_fraction": row.precursor_purity_fraction,
            "pure_precursor_mass_g": row.pure_precursor_mass_g,
            "commercial_precursor_mass_g": row.commercial_precursor_mass_g,
            "precursor_amount_mol": row.precursor_amount_mol,
            "concentration_commercial_g_ml": row.concentration_commercial_g_ml,
            "concentration_pure_mol_l": row.concentration_pure_mol_l,
        }
        for row in result.requirements
    ]


def _write_xlsx(result: RecipeResult, path: Path) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            'XLSX export requires the optional dependency: pip install "impregcalc[xlsx]"'
        ) from exc
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Recipe"
    summary.append(["Field", "Value", "Unit or convention"])
    summary.append(["Support", result.request.support_name, ""])
    summary.append(["Support mass", result.request.support_mass_g, "g"])
    summary.append(["Method", result.request.method.value, ""])
    summary.append(["Pore volume", result.request.pore_volume_ml_g, "mL g-1"])
    summary.append(["Nominal pore-volume product", result.nominal_pore_volume_ml, "mL"])
    summary.append(["Selected solution volume", result.selected_solution_volume_ml, "mL"])
    summary.append(["Final formulated mass", result.final_formulated_mass_g, "g"])
    components = workbook.create_sheet("Components")
    rows = _csv_rows(result)
    components.append(list(rows[0]))
    for component_row in rows:
        components.append(list(component_row.values()))
    cycles = workbook.create_sheet("Cycles")
    cycles.append(["Component", "Cycle", "Fraction", "Precursor / g", "Solution / mL"])
    for requirement in result.requirements:
        for cycle in requirement.cycle_portions:
            cycles.append(
                [
                    requirement.name,
                    cycle.cycle,
                    cycle.fraction,
                    cycle.precursor_mass_g,
                    cycle.solution_volume_ml,
                ]
            )
    notes = workbook.create_sheet("Assumptions and warnings")
    notes.append(["Type", "Text"])
    for item in result.assumptions:
        notes.append(["assumption", item])
    for item in result.warnings:
        notes.append(["warning", item])
    workbook.save(path)


def write_result(result: RecipeResult, path: str | Path) -> Path:
    """Write JSON, Markdown, CSV, or XLSX according to the file suffix."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    suffix = destination.suffix.lower()
    if suffix == ".json":
        destination.write_text(
            json.dumps(result.as_dict(), indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
    elif suffix in {".md", ".markdown", ".txt"}:
        destination.write_text(format_markdown(result), encoding="utf-8")
    elif suffix == ".csv":
        rows = _csv_rows(result)
        with destination.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    elif suffix == ".xlsx":
        _write_xlsx(result, destination)
    else:
        raise ValueError("Output suffix must be .json, .md, .txt, .csv, or .xlsx")
    return destination
